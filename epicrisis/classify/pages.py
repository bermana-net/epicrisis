"""Pages to classify, derived from inventory records, and the payload sent for each page.

A payload is either the page text or a PNG rendered for that page alone. It never carries a
file name, folder name, path, folder year or page number, and images are re-encoded without
metadata (phone photos keep GPS coordinates and device names in EXIF).
"""

import hashlib
from datetime import datetime
import io
import threading
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
import openpyxl
import xlrd
import pypdfium2
from PIL import Image, ImageOps
from pypdf import PdfReader

from epicrisis.inventory import legacy
from epicrisis.inventory.probes import MIN_TEXT_CHARS_PER_PAGE

PDFIUM_LOCK = threading.Lock()
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MAX_IMAGE_EDGE = 1600
MAX_RENDER_DPI = 200
# Second pass for small print: the page is rendered larger and cut into four overlapping
# quarters, each sent at MAX_IMAGE_EDGE, which gives the model about twice the detail.
ZOOM_RENDER_EDGE = 3200
ZOOM_RENDER_DPI = 300
CLOSE_UP_SHARE = 0.55
MAX_TEXT_CHARS = 20_000
MAX_SHEET_ROWS = 200


@dataclass(frozen=True)
class PageRef:
    file_sha256: str
    page: int  # 1-based position among all pages of the file
    route: str  # "text" or "vision"
    part: str  # "pdf", "image", "office_text" or "office_image"
    index: int  # 0-based position within its part
    record: dict = field(compare=False, hash=False, repr=False)


@dataclass
class Payload:
    text: str | None = None
    image_path: Path | None = None
    close_ups: list[Path] = field(default_factory=list)


class PageUnreadable(Exception):
    """The page cannot be turned into text or an image. Recorded without calling a model."""


def page_refs(record: dict) -> list[PageRef]:
    if record.get("skipped") or "error" in record or "unsupported" in record or "sha256" not in record:
        return []
    refs: list[PageRef] = []
    counts: dict[str, int] = {}

    def add(route: str, part: str) -> None:
        refs.append(PageRef(record["sha256"], len(refs) + 1, route, part, counts.get(part, 0), record))
        counts[part] = counts.get(part, 0) + 1

    kind = record.get("category")
    if kind == "pdf":
        garbled = set(record["pdf"].get("garbled_text_pages", []))
        for number, chars in enumerate(record["pdf"].get("text_chars_per_page", []), 1):
            readable = chars >= MIN_TEXT_CHARS_PER_PAGE and number not in garbled
            add("text" if readable else "vision", "pdf")
    elif kind == "image":
        for _ in range(record["image"]["frames"]):
            add("vision", "image")
    elif kind in ("word", "excel"):
        facts = record[kind]
        if kind == "excel":
            for _ in range(facts["sheets"]):
                add("text", "office_text")
        elif facts["text_chars"]:
            add("text", "office_text")
        for _ in range(facts.get("embedded_images", 0)):
            add("vision", "office_image")
    return refs


def materialize(ref: PageRef, archive_root: Path, workdir: Path) -> Payload:
    data = _file_bytes(ref.record, archive_root)
    try:
        if ref.part == "pdf":
            if ref.route == "text":
                return Payload(text=_pdf_text(data, ref.index))
            return _png(_render_pdf_page(data, ref.index), ref, workdir)
        if ref.part == "image":
            with Image.open(io.BytesIO(data)) as image:
                image.seek(ref.index)
                return _png(image, ref, workdir)
        kind = ref.record["category"]
        if ref.part == "office_text":
            text = _word_text(data) if kind == "word" else _sheet_text(data, ref.index)
            return Payload(text=text[:MAX_TEXT_CHARS])
        name = _office_media_names(data, kind)[ref.index]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            with Image.open(io.BytesIO(archive.read(name))) as image:
                return _png(image, ref, workdir)
    except Exception as exc:
        # Only the exception type is kept: messages can quote document content.
        raise PageUnreadable(type(exc).__name__) from exc


def _file_bytes(record: dict, archive_root: Path) -> bytes:
    data = (archive_root / record["path"]).read_bytes()
    if hashlib.sha256(data).hexdigest() != record["sha256"]:
        raise PageUnreadable("file changed since inventory")
    wrapper = record.get("wrapper")
    data = data[wrapper["payload_offset"] :] if wrapper else data
    if record.get("word", {}).get("format") == "doc":
        # Word 97-2003 is read from a .docx copy, so text and pictures come out as for .docx.
        data = legacy.doc_to_docx(data)
    return data


def _pdf_text(data: bytes, index: int) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt("")
    return (reader.pages[index].extract_text() or "")[:MAX_TEXT_CHARS]


def _render_pdf_page(data: bytes, index: int, max_edge: int = MAX_IMAGE_EDGE, max_dpi: int = MAX_RENDER_DPI) -> Image.Image:
    # PDFium is not thread-safe: parallel runs and the dashboard render one page at a time.
    with PDFIUM_LOCK:
        return _render_pdf_page_unlocked(data, index, max_edge, max_dpi)


def _render_pdf_page_unlocked(data: bytes, index: int, max_edge: int, max_dpi: int) -> Image.Image:
    document = pypdfium2.PdfDocument(data)
    try:
        page = document[index]
        try:
            width, height = page.get_size()
            scale = min(max_edge / max(width, height), max_dpi / 72)
            return page.render(scale=scale).to_pil().copy()
        finally:
            page.close()
    finally:
        document.close()


def _png(image: Image.Image, ref: PageRef, workdir: Path) -> Payload:
    path = workdir / f"{ref.file_sha256[:16]}.png"
    _clean_image(image).save(path, "PNG")
    return Payload(image_path=path)


def _clean_image(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    image.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
    # A fresh image carries pixels only: no EXIF, ICC profile or text chunks.
    clean = Image.new(image.mode, image.size)
    clean.paste(image)
    return clean


def original_png(ref: PageRef, archive_root: Path) -> bytes:
    """The page as a PNG for viewing on this server. Nothing here leaves the machine."""
    data = _file_bytes(ref.record, archive_root)
    try:
        buffer = io.BytesIO()
        _clean_image(_page_image(data, ref)).save(buffer, "PNG")
    except PageUnreadable:
        raise
    except Exception as exc:
        raise PageUnreadable(type(exc).__name__) from exc
    return buffer.getvalue()


def document_payloads(refs: list[PageRef], archive_root: Path, workdir: Path, zoom: bool = False, always_images: bool = False) -> list[Payload]:
    """Payloads for the pages of one document, in order.

    If any page is a scan, PDF pages all go as images, so the model reads the document one way.
    Word and Excel text pages have no image and always go as text. Image files are named by
    the file hash and the position in this document, never by the page number in the file.
    With zoom, every image page also gets four overlapping close-ups. always_images renders the
    PDF pages even where a text layer exists: a text layer holds no axis labels, no stamp and no
    small print outside the text flow, so a second reading of such a page has to look at it.
    """
    data = _file_bytes(refs[0].record, archive_root)
    as_images = always_images or any(ref.route == "vision" for ref in refs)
    payloads = []
    try:
        for position, ref in enumerate(refs, 1):
            if ref.part == "office_text":
                kind = ref.record["category"]
                text = _word_text(data) if kind == "word" else _sheet_text(data, ref.index)
                payloads.append(Payload(text=text[:MAX_TEXT_CHARS]))
            elif ref.part == "pdf" and not as_images:
                payloads.append(Payload(text=_pdf_text(data, ref.index)))
            else:
                stem = f"{ref.file_sha256[:16]}-{position:02d}"
                image = _page_image(data, ref, zoom)
                path = workdir / f"{stem}.png"
                _clean_image(image).save(path, "PNG")
                close_ups = []
                if zoom:
                    for part, close_up in zip("abcd", _close_ups(image), strict=False):
                        close_ups.append(workdir / f"{stem}{part}.png")
                        _clean_image(close_up).save(close_ups[-1], "PNG")
                payloads.append(Payload(image_path=path, close_ups=close_ups))
    except PageUnreadable:
        raise
    except Exception as exc:
        raise PageUnreadable(type(exc).__name__) from exc
    return payloads


def _close_ups(image: Image.Image) -> list[Image.Image]:
    """Four overlapping quarters of the upright page, row by row."""
    image = ImageOps.exif_transpose(image)
    width, height = image.size
    cut_x, cut_y = int(width * CLOSE_UP_SHARE), int(height * CLOSE_UP_SHARE)
    return [
        image.crop((left, top, left + cut_x, top + cut_y))
        for top in (0, height - cut_y)
        for left in (0, width - cut_x)
    ]


def _page_image(data: bytes, ref: PageRef, zoom: bool = False) -> Image.Image:
    if ref.part == "pdf":
        if zoom:
            return _render_pdf_page(data, ref.index, ZOOM_RENDER_EDGE, ZOOM_RENDER_DPI)
        return _render_pdf_page(data, ref.index)
    if ref.part == "image":
        with Image.open(io.BytesIO(data)) as source:
            source.seek(ref.index)
            return source.copy()
    if ref.part == "office_image":
        name = _office_media_names(data, ref.record["category"])[ref.index]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            with Image.open(io.BytesIO(archive.read(name))) as source:
                return source.copy()
    raise PageUnreadable("text page without an image")


def _word_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _sheet_text(data: bytes, index: int) -> str:
    if data.startswith(OLE_SIGNATURE):
        return _legacy_sheet_text(data, index)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[index]
        return "\n".join(
            "\t".join("" if value is None else str(value) for value in row)
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True)
        )
    finally:
        workbook.close()


def _legacy_sheet_text(data: bytes, index: int) -> str:
    """An Excel 97-2003 sheet as tab-separated text, dates and whole numbers as a person sees them."""
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(index)
        lines = []
        for row in range(min(sheet.nrows, MAX_SHEET_ROWS)):
            cells = []
            for cell in sheet.row(row):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    moment = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                    cells.append(moment.strftime("%d.%m.%Y") if moment.time() == datetime.min.time() else moment.strftime("%d.%m.%Y %H:%M"))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    cells.append(str(int(cell.value)))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    cells.append("")
                else:
                    cells.append(str(cell.value))
            lines.append("\t".join(cells))
        return "\n".join(lines)
    finally:
        workbook.release_resources()


def _office_media_names(data: bytes, kind: str) -> list[str]:
    if data.startswith(OLE_SIGNATURE):
        return []  # Excel 97-2003: pictures are not counted in inventory either
    prefix = "word/media/" if kind == "word" else "xl/media/"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return sorted(name for name in archive.namelist() if name.startswith(prefix) and not name.endswith("/"))
