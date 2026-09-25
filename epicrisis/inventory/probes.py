"""File type detection and per-format probes.

A probe reads a file, or the payload inside a saved HTTP response, without modifying it and
returns format-specific facts. It raises UnsupportedFormat for files it cannot parse by
design, and any other exception for files that are damaged.
"""

import logging
import re
import unicodedata
import zipfile
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import docx
import magic
import openpyxl
import xlrd
from PIL import Image, UnidentifiedImageError
from pypdf import PasswordType, PdfReader

# A page with fewer visible characters than this is treated as having no text layer.
# Scans carry stamps and page numbers as text: the real archive showed scanned pages with
# 1-19 characters, real text pages with 500 or more, and nothing in between.
MIN_TEXT_CHARS_PER_PAGE = 100

# A text layer can look fine on screen and still extract as garbage: fonts with a private
# encoding come out as control characters and Latin Extended-B. Such pages go as images.
# In the real archive broken pages had 45-63% of these characters, readable ones 1.1% at most.
MAX_GARBLED_SHARE = 0.1

HEAD_BYTES = 64 * 1024
HEADER_LINE = re.compile(rb"^[ \t]*[A-Za-z0-9-]+:[^\n]*$")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LEGACY_OFFICE_MIMES = {
    "application/msword",
    "application/vnd.ms-excel",
    "application/vnd.ms-office",
    "application/x-ole-storage",
    "application/CDFV2",
}

logging.getLogger("pypdf").setLevel(logging.ERROR)

Source = Path | BinaryIO


class UnsupportedFormat(Exception):
    pass


def http_payload(path: Path) -> tuple[int, BytesIO] | None:
    """Offset and body of a file saved as a raw HTTP response, or None for ordinary files.

    Browsers and download tools sometimes store the status line and headers in front of a PDF.
    Header lines may be indented and the blank line before the body may be missing, so the
    header block ends at the first line that is blank or does not look like a header.
    """
    with path.open("rb") as fh:
        head = fh.read(HEAD_BYTES)
    if not head.startswith(b"HTTP/1."):
        return None

    headers = []
    offset = head.find(b"\n") + 1
    while True:
        end = head.find(b"\n", offset)
        if offset == 0 or end == -1:
            raise UnsupportedFormat("saved HTTP response without a complete header")
        line = head[offset:end]
        if not line.strip():
            offset = end + 1
            break
        if not HEADER_LINE.match(line):
            break
        headers.append(line.decode("latin-1").strip().lower())
        offset = end + 1
    while offset < len(head) and head[offset] in b" \t\r\n":
        offset += 1

    if any(header.startswith(("transfer-encoding: chunked", "content-encoding:")) for header in headers):
        raise UnsupportedFormat("saved HTTP response with an encoded body")
    return offset, BytesIO(path.read_bytes()[offset:])


def detect_mime(source: Source) -> str:
    """MIME type from the content signature, never from the extension."""
    if isinstance(source, Path):
        mime = magic.from_file(str(source), mime=True)
    else:
        source.seek(0)
        mime = magic.from_buffer(source.read(HEAD_BYTES), mime=True)
        source.seek(0)
    if mime in ("application/zip", "application/octet-stream"):
        mime = _sniff_ooxml(source) or mime
    return mime


def _sniff_ooxml(source: Source) -> str | None:
    # Some libmagic builds report .docx/.xlsx as a plain zip archive.
    try:
        names = _zip_names(source)
    except (zipfile.BadZipFile, OSError):
        return None
    if "word/document.xml" in names:
        return DOCX_MIME
    if "xl/workbook.xml" in names:
        return XLSX_MIME
    return None


def _zip_names(source: Source) -> list[str]:
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
    if not isinstance(source, Path):
        source.seek(0)
    return names


def category(mime: str) -> str:
    if mime == "application/pdf":
        return "pdf"
    if mime.startswith("image/"):
        return "image"
    if mime == DOCX_MIME:
        return "word"
    if mime == XLSX_MIME:
        return "excel"
    if mime in LEGACY_OFFICE_MIMES:
        return "legacy_office"
    if mime == "inode/x-empty":
        return "empty"
    if mime.startswith("text/"):
        return "text"
    return "other"


def _visible_chars(text: str) -> int:
    return sum(1 for ch in text if not ch.isspace())


def _garbled_char(ch: str) -> bool:
    code = ord(ch)
    # Latin Extended-B, except the Romanian comma-below letters that real text uses.
    if 0x0180 <= code <= 0x024F and not 0x0218 <= code <= 0x021B:
        return True
    return code == 0xFFFD or unicodedata.category(ch) in ("Cc", "Co", "Cn", "Cs")


def text_looks_garbled(text: str) -> bool:
    visible = [ch for ch in text if not ch.isspace()]
    return bool(visible) and sum(map(_garbled_char, visible)) / len(visible) > MAX_GARBLED_SHARE


def probe_pdf(source: Source, mime: str) -> dict:
    reader = PdfReader(source)
    encrypted = reader.is_encrypted
    if encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
        return {"encrypted": True}

    chars_per_page = []
    garbled_pages = []
    has_images = False
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        chars_per_page.append(_visible_chars(text))
        if chars_per_page[-1] >= MIN_TEXT_CHARS_PER_PAGE and text_looks_garbled(text):
            garbled_pages.append(number)
        has_images = has_images or _has_images(page.get("/Resources"))

    pages = len(chars_per_page)
    pages_with_text = sum(1 for chars in chars_per_page if chars >= MIN_TEXT_CHARS_PER_PAGE) - len(garbled_pages)
    if pages and pages_with_text == pages:
        text_layer = "full"
    elif pages_with_text:
        text_layer = "partial"
    else:
        text_layer = "none"

    return {
        "encrypted": encrypted,
        "pages": pages,
        "pages_with_text": pages_with_text,
        "text_layer": text_layer,
        "text_chars_per_page": chars_per_page,
        "garbled_text_pages": garbled_pages,
        "has_images": has_images,
    }


def _has_images(resources, depth: int = 0) -> bool:
    # Images sit in the page's XObject resources, possibly nested inside Form XObjects.
    if resources is None or depth > 5:
        return False
    xobjects = resources.get_object().get("/XObject")
    if xobjects is None:
        return False
    for ref in xobjects.get_object().values():
        xobject = ref.get_object()
        subtype = xobject.get("/Subtype")
        if subtype == "/Image":
            return True
        if subtype == "/Form" and _has_images(xobject.get("/Resources"), depth + 1):
            return True
    return False


def probe_image(source: Source, mime: str) -> dict:
    try:
        with Image.open(source) as image:
            dpi = image.info.get("dpi")
            facts = {
                "format": image.format,
                "width": image.width,
                "height": image.height,
                "dpi": [round(float(value)) for value in dpi] if dpi else None,
                "frames": getattr(image, "n_frames", 1),
            }
            # Decoding the pixels is the only reliable way to catch truncated files.
            image.load()
    except UnidentifiedImageError as exc:
        # Pillow knows the format but cannot read this file: damaged. Otherwise: unsupported.
        Image.init()
        if mime in Image.MIME.values():
            raise
        raise UnsupportedFormat("image format not supported by Pillow") from exc
    return facts


def probe_excel(source: Source, mime: str) -> dict:
    # Pictures pasted into a workbook are often scans and need the vision pass.
    embedded_images = sum(1 for name in _zip_names(source) if name.startswith("xl/media/"))
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        dimensions = []
        for sheet in workbook.worksheets:
            rows, columns = sheet.max_row, sheet.max_column
            if rows is None or columns is None:
                # The sheet has no stored dimension; count by reading it.
                rows = columns = 0
                for row in sheet.iter_rows(values_only=True):
                    rows += 1
                    columns = max(columns, len(row))
            dimensions.append({"rows": rows, "columns": columns})
    finally:
        workbook.close()
    return {"sheets": len(dimensions), "sheet_dimensions": dimensions, "embedded_images": embedded_images}


def probe_legacy_excel(source: Source, mime: str) -> dict:
    """Excel 97-2003 workbooks. Raises UnsupportedFormat for other legacy Office files (.doc)."""
    try:
        data = source.read_bytes() if isinstance(source, Path) else _read_all(source)
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        raise UnsupportedFormat("legacy Office file that is not an Excel workbook") from exc
    try:
        dimensions = [{"rows": sheet.nrows, "columns": sheet.ncols} for sheet in (workbook.sheet_by_index(i) for i in range(workbook.nsheets))]
    finally:
        workbook.release_resources()
    return {"sheets": len(dimensions), "sheet_dimensions": dimensions, "embedded_images": 0, "format": "xls"}


def _read_all(source: BinaryIO) -> bytes:
    source.seek(0)
    data = source.read()
    source.seek(0)
    return data


def probe_word(source: Source, mime: str) -> dict:
    # Pictures pasted into a document are often scans and need the vision pass.
    embedded_images = sum(1 for name in _zip_names(source) if name.startswith("word/media/"))
    document = docx.Document(str(source) if isinstance(source, Path) else source)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return {"text_chars": sum(_visible_chars(text) for text in texts), "embedded_images": embedded_images}
