"""A spreadsheet, new or old: one page per sheet, as tab-separated text, and its pictures.

A sheet is a page because that is how a person reads a workbook — one tab at a time — and because
a model asked for a whole workbook at once answers about the wrong sheet. Dates and whole numbers
of an old workbook are written as a person sees them, not as the floating-point numbers they are
stored as. See epicrisis/readers/__init__.py for what every reader of a kind of file answers.
"""

import io
from pathlib import Path
from datetime import datetime

import openpyxl
import xlrd
from PIL import Image

from epicrisis.inventory.probes import MAX_TEXT_CHARS, Source, UnsupportedFormat, _read_all, _zip_names
from epicrisis.readers import office
from epicrisis.readers.office import OLE_SIGNATURE

MAX_SHEET_ROWS = 200

def probe(source: Source, mime: str) -> dict:
    # Pictures pasted into a workbook are often scans and need the vision pass. Which entries of
    # the zip are pictures is office.py's answer, for the same reason it is in word.py: a zip may
    # hold the directory "xl/media/" as an entry of its own, and counted as a picture it gave the
    # workbook a page that nothing could be read for.
    embedded_images = len(office.pictures_among(_zip_names(source), "excel"))
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        dimensions = []
        for sheet in workbook.worksheets:
            rows, columns = sheet.max_row, sheet.max_column
            if rows is None or columns is None:
                # The sheet has no stored dimension; count by reading it.
                rows = columns = 0
                for row in sheet.iter_rows(values_only=True):
                    rows += 1
                    columns = max(columns, len(row))
            dimensions.append({"rows": rows, "columns": columns})
    finally:
        workbook.close()
    return {"sheets": len(dimensions), "sheet_dimensions": dimensions, "embedded_images": embedded_images}

def probe_legacy_excel(source: Source, mime: str) -> dict:
    """Excel 97-2003 workbooks. Raises UnsupportedFormat for other legacy Office files (.doc)."""
    try:
        data = source.read_bytes() if isinstance(source, Path) else _read_all(source)
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        raise UnsupportedFormat("legacy Office file that is not an Excel workbook") from exc
    try:
        dimensions = [{"rows": sheet.nrows, "columns": sheet.ncols} for sheet in (workbook.sheet_by_index(i) for i in range(workbook.nsheets))]
    finally:
        workbook.release_resources()
    return {"sheets": len(dimensions), "sheet_dimensions": dimensions, "embedded_images": 0, "format": "xls"}

def _sheet_text(data: bytes, index: int) -> str:
    if data.startswith(OLE_SIGNATURE):
        return _legacy_sheet_text(data, index)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[index]
        return "\n".join(
            "\t".join("" if value is None else str(value) for value in row)
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True)
        )
    finally:
        workbook.close()

def _legacy_sheet_text(data: bytes, index: int) -> str:
    """An Excel 97-2003 sheet as tab-separated text, dates and whole numbers as a person sees them."""
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(index)
        lines = []
        for row in range(min(sheet.nrows, MAX_SHEET_ROWS)):
            cells = []
            for cell in sheet.row(row):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    moment = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                    cells.append(moment.strftime("%d.%m.%Y") if moment.time() == datetime.min.time() else moment.strftime("%d.%m.%Y %H:%M"))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    cells.append(str(int(cell.value)))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    cells.append("")
                else:
                    cells.append(str(cell.value))
            lines.append("\t".join(cells))
        return "\n".join(lines)
    finally:
        workbook.release_resources()


def pages(record: dict) -> list[tuple[str, str, int | None]]:
    """One page per sheet, then one for every picture pasted into the workbook."""
    facts = record["excel"]
    out = [("text", "office_text", None) for _ in range(facts["sheets"])]
    return out + [("vision", "office_image", None) for _ in range(facts.get("embedded_images", 0))]


def text_of(data: bytes, ref) -> str:
    return _sheet_text(data, ref.index)[:MAX_TEXT_CHARS]


def image_of(data: bytes, ref, zoom: bool = False) -> Image.Image:
    return office.image_of(data, ref, zoom)
