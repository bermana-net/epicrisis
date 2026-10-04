"""A Word document, new or old: text somebody typed, and sometimes a scan pasted into it.

Its text is one page however long it is — a document is not paginated until it is printed — and
every picture inside it is a page of its own, because in an archive like this a picture pasted
into a letter is usually a photograph of a form. A Word 97-2003 file is read from a .docx copy
made by LibreOffice, so text and pictures come out the same way. See
epicrisis/readers/__init__.py for what every reader of a kind of file answers.
"""

import io
from pathlib import Path

import docx
from PIL import Image

from epicrisis.inventory.probes import MAX_TEXT_CHARS, Source, _visible_chars, _zip_names
from epicrisis.readers import office

def probe(source: Source, mime: str) -> dict:
    # Pictures pasted into a document are often scans and need the vision pass. Which entries of
    # the zip are pictures is office.py's answer and not a second reading of the same names here:
    # counted with a prefix alone, the directory entry "word/media/" that python's zipfile and
    # LibreOffice both write was a picture, and the file was given a page with nothing behind it.
    embedded_images = len(office.pictures_among(_zip_names(source), "word"))
    document = docx.Document(str(source) if isinstance(source, Path) else source)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return {"text_chars": sum(_visible_chars(text) for text in texts), "embedded_images": embedded_images}

def _word_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def pages(record: dict) -> list[tuple[str, str, int | None]]:
    """The text, if it has any, and then one page for every picture pasted into it."""
    facts = record["word"]
    out = [("text", "office_text", None)] if facts["text_chars"] else []
    return out + [("vision", "office_image", None) for _ in range(facts.get("embedded_images", 0))]


def text_of(data: bytes, ref) -> str:
    return _word_text(data)[:MAX_TEXT_CHARS]


def image_of(data: bytes, ref, zoom: bool = False) -> Image.Image:
    return office.image_of(data, ref, zoom)
