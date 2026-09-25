"""Classify tests. Synthetic files only; the model is replaced by a fake backend."""

import json
import os
import subprocess
from pathlib import Path

import docx
import pytest
from PIL import Image
from typer.testing import CliRunner

from epicrisis.classify import backend as backend_module
from epicrisis.classify.backend import (
    PROMPT_VERSION,
    BackendError,
    Classification,
    ClaudeCodeBackend,
    UsageLimitReached,
    parse_result,
)
from epicrisis.classify.pages import Payload, materialize, page_refs
from epicrisis.classify.report import group_documents, latest_pages, render_summary
from epicrisis.classify.run import classify_source, is_running, parse_years, sample_refs
from epicrisis.cli import app
from epicrisis.consent import record_consent
from epicrisis.inventory.run import write_inventory
from epicrisis.records import read_records
from epicrisis.sources import SourceRegistry, source_output_dir
from test_inventory import SYNTHETIC_TEXT, make_scan_pdf, make_text_pdf

# A folder name that must never reach the model, the report or the CLI output.
PRIVATE_NAME = "Clinic Ivanova oncology"

FIELDS = {
    "doc_type": "lab_panel",
    "page_role": "first",
    "language": "en",
    "date_on_page": "12.03.2019",
    "provider_on_page": "Synthetic Lab",
    "has_tabular_results": True,
    "legible": True,
    "confidence": 0.95,
}


class FakeBackend:
    name = "fake"
    model = "fake-model-1"

    def __init__(self, fail_from: int | None = None, limit_from: int | None = None, model: str | None = None,
                 **rest):  # a backend is handed the call its engine makes; this one carries its own
        if model:
            self.model = model
        self.calls = []
        self.fail_from = fail_from
        self.limit_from = limit_from

    def classify(self, payload: Payload, workdir: Path) -> Classification:
        if self.limit_from is not None and len(self.calls) >= self.limit_from:
            raise UsageLimitReached("usage limit reached")
        call = {"text": payload.text, "image_name": None, "workdir": workdir}
        if payload.image_path is not None:
            with Image.open(payload.image_path) as image:
                call.update(image_name=payload.image_path.name, info=dict(image.info), exif=len(image.getexif()), size=image.size)
        self.calls.append(call)
        if self.fail_from is not None and len(self.calls) > self.fail_from:
            raise BackendError("model call failed")
        return Classification(fields=dict(FIELDS), model="fake-model-1-20260101")


@pytest.fixture
def setup(tmp_path):
    archive = tmp_path / PRIVATE_NAME
    (archive / "2003").mkdir(parents=True)
    (archive / "2019").mkdir()
    make_scan_pdf(archive / "2003" / "scan_a.pdf", pages=2)
    make_scan_pdf(archive / "2003" / "scan_b.pdf")
    make_text_pdf(archive / "2019" / "labs_a.pdf", [SYNTHETIC_TEXT, SYNTHETIC_TEXT + " continued"])
    make_text_pdf(archive / "2019" / "labs_b.pdf", [SYNTHETIC_TEXT + " second laboratory"])
    exif = Image.Exif()
    exif[0x010F] = "SyntheticMaker"
    exif[0x0110] = "SyntheticPhone"
    Image.new("RGB", (2400, 1800), "white").save(archive / "2019" / "photo.jpg", "JPEG", exif=exif.tobytes())

    data_dir = tmp_path / "data"
    source = SourceRegistry(data_dir).add(str(archive))
    output = source_output_dir(data_dir.resolve(), source.id)
    write_inventory(archive, output / "inventory.jsonl")
    return data_dir.resolve(), source, archive, output


def records_by_name(output: Path) -> dict[str, dict]:
    return {record["name"]: record for record in read_records(output / "inventory.jsonl")}


def test_page_routes(setup):
    _, _, _, output = setup
    records = records_by_name(output)
    assert [ref.route for ref in page_refs(records["scan_a.pdf"])] == ["vision", "vision"]
    assert [ref.route for ref in page_refs(records["labs_a.pdf"])] == ["text", "text"]
    assert [(ref.route, ref.part) for ref in page_refs(records["photo.jpg"])] == [("vision", "image")]
    assert page_refs({"sha256": "x", "category": "legacy_office"}) == []
    assert page_refs({"sha256": "x", "category": "pdf", "error": "PdfReadError"}) == []


def test_word_pages_text_then_images(tmp_path):
    picture = tmp_path / "scan.png"
    Image.new("RGB", (60, 40), "white").save(picture)
    document = docx.Document()
    document.add_paragraph("Synthetic cover note")
    document.add_picture(str(picture))
    root = tmp_path / "archive"
    root.mkdir()
    document.save(root / "note.docx")
    write_inventory(root, tmp_path / "inventory.jsonl")
    record = next(read_records(tmp_path / "inventory.jsonl"))

    refs = page_refs(record)

    assert [(ref.page, ref.route, ref.part, ref.index) for ref in refs] == [
        (1, "text", "office_text", 0),
        (2, "vision", "office_image", 0),
    ]
    workdir = tmp_path / "work"
    workdir.mkdir()
    assert "Synthetic cover note" in materialize(refs[0], root, workdir).text
    assert materialize(refs[1], root, workdir).image_path.name == f"{record['sha256'][:16]}.png"


def test_payloads_carry_no_names_or_metadata(setup):
    data_dir, source, archive, _ = setup
    fake = FakeBackend()

    classify_source(data_dir, source, fake)

    assert len(fake.calls) == 7
    for call in fake.calls:
        if call["text"] is not None:
            assert PRIVATE_NAME not in call["text"] and "labs_a" not in call["text"]
            assert "Synthetic hemoglobin" in call["text"]
        else:
            assert len(call["image_name"]) == len("0123456789abcdef.png")
            assert call["exif"] == 0 and "exif" not in call["info"] and "icc_profile" not in call["info"]
            assert max(call["size"]) <= 1600
        assert not call["workdir"].exists()
        assert not call["workdir"].is_relative_to(data_dir) and not call["workdir"].is_relative_to(archive)


def test_lock_is_held_during_run_and_blocks_second_cli_run(setup, monkeypatch):
    data_dir, source, _, output = setup
    seen = []

    class LockWatcher(FakeBackend):
        def classify(self, payload, workdir):
            seen.append(is_running(output))
            return super().classify(payload, workdir)

    classify_source(data_dir, source, LockWatcher(), limit=1)
    assert seen == [True] and not is_running(output)

    (output / "classify.lock").write_text(json.dumps({"pid": os.getpid()}))
    monkeypatch.setattr(backend_module, "ClaudeCodeBackend", FakeBackend)
    record_consent(data_dir, FakeBackend.name)
    result = CliRunner().invoke(app, ["classify", "--data-dir", str(data_dir)])
    assert result.exit_code == 2 and "already in progress" in result.output


def test_results_provenance_and_resume(setup):
    data_dir, source, _, output = setup

    first = classify_source(data_dir, source, FakeBackend())
    second_backend = FakeBackend()
    second = classify_source(data_dir, source, second_backend)

    assert (first.total, first.classified) == (7, 7)
    assert (second.classified, second.already_done, len(second_backend.calls)) == (0, 7, 0)
    lines = list(read_records(output / "classify.jsonl"))
    assert len(lines) == 7
    assert lines[0]["provenance"]["model"] == "fake-model-1-20260101"
    assert lines[0]["provenance"]["prompt_version"] == PROMPT_VERSION
    assert {line["route"] for line in lines} == {"text", "vision"}


def test_page_is_classified_again_when_its_route_changes(setup):
    data_dir, source, _, output = setup
    classify_source(data_dir, source, FakeBackend())

    records = list(read_records(output / "inventory.jsonl"))
    labs = next(record for record in records if record.get("name") == "labs_a.pdf")
    labs["pdf"]["garbled_text_pages"] = [2]
    (output / "inventory.jsonl").write_text("".join(json.dumps(record) + "\n" for record in records))
    again = FakeBackend()
    stats = classify_source(data_dir, source, again)

    assert (stats.classified, stats.already_done) == (1, 6)
    page = next(p for p in latest_pages(output / "classify.jsonl") if (p["file_sha256"], p["page"]) == (labs["sha256"], 2))
    assert page["route"] == "vision"


def test_usage_limit_stops_and_resumes(setup):
    data_dir, source, _, _ = setup

    stopped = classify_source(data_dir, source, FakeBackend(limit_from=3), workers=1)
    resumed = classify_source(data_dir, source, FakeBackend())

    assert (stopped.stopped, stopped.classified) == ("usage_limit", 3)
    assert (resumed.classified, resumed.already_done) == (4, 3)


def test_failed_pages_are_retried(setup):
    data_dir, source, _, _ = setup

    partly = classify_source(data_dir, source, FakeBackend(fail_from=5), workers=1)
    retry = classify_source(data_dir, source, FakeBackend())

    assert (partly.classified, partly.failed) == (5, 2)
    assert (retry.classified, retry.already_done) == (2, 5)


def test_limit_and_sample(setup):
    data_dir, source, _, output = setup
    records = list(read_records(output / "inventory.jsonl"))

    sample = sample_refs(records)

    assert sorted(ref.route for ref in sample) == ["text", "text", "vision", "vision", "vision"]
    assert all(ref.page == 1 for ref in sample)
    assert len({ref.file_sha256 for ref in sample}) == 5
    assert classify_source(data_dir, source, FakeBackend(), limit=2).classified == 2


@pytest.mark.parametrize(
    ("spec", "expected"),
    [("2003", {2003}), ("1992-1994", {1992, 1993, 1994}), (" 2001, 2019-2020 ", {2001, 2019, 2020})],
)
def test_parse_years(spec, expected):
    assert parse_years(spec) == expected


@pytest.mark.parametrize("spec", ["", "20o3", "2005-2001", "1992-2003-2005"])
def test_parse_years_rejects_bad_input(spec):
    with pytest.raises(ValueError):
        parse_years(spec)


def test_years_filter_limits_run_to_those_folders(setup, monkeypatch):
    data_dir, source, _, output = setup
    fake = FakeBackend()

    stats = classify_source(data_dir, source, fake, years={2003})

    assert (stats.total, stats.classified) == (3, 3)
    classified_files = {line["file_sha256"] for line in read_records(output / "classify.jsonl")}
    by_name = records_by_name(output)
    assert classified_files == {by_name["scan_a.pdf"]["sha256"], by_name["scan_b.pdf"]["sha256"]}

    monkeypatch.setattr(backend_module, "ClaudeCodeBackend", FakeBackend)
    record_consent(data_dir, FakeBackend.name)
    runner = CliRunner()
    result = runner.invoke(app, ["classify", "--data-dir", str(data_dir), "--years", "2019"])
    bad = runner.invoke(app, ["classify", "--data-dir", str(data_dir), "--years", "2019-2003"])
    assert result.exit_code == 0 and "This run: classified 4" in result.output
    assert bad.exit_code == 2 and "backwards" in bad.output


def test_changed_file_is_unreadable_without_model_call(setup):
    data_dir, source, archive, output = setup
    make_text_pdf(archive / "2019" / "labs_b.pdf", [SYNTHETIC_TEXT + " edited after inventory"])
    fake = FakeBackend()

    stats = classify_source(data_dir, source, fake)

    assert (stats.unreadable, len(fake.calls)) == (1, 6)
    errors = [line for line in read_records(output / "classify.jsonl") if "error" in line]
    assert errors[0]["error"] == "file changed since inventory"


def test_documents_and_summary_hold_no_content(setup):
    data_dir, source, _, output = setup
    classify_source(data_dir, source, FakeBackend())
    pages = latest_pages(output / "classify.jsonl")

    summary = render_summary(pages, total_pages=7)

    assert "Pages classified: 7 of 7" in summary
    assert "Documents: 7" in summary
    assert "12.03.2019" not in summary and "Synthetic Lab" not in summary


def test_group_documents_splits_on_first_pages():
    pages = [
        {"file_sha256": "a", "page": 1, "page_role": "first"},
        {"file_sha256": "a", "page": 2, "page_role": "continuation"},
        {"file_sha256": "a", "page": 3, "page_role": "first"},
        {"file_sha256": "b", "page": 1, "page_role": "continuation"},
        {"file_sha256": "b", "page": 2, "error": "PdfReadError"},
    ]
    assert [len(document) for document in group_documents(pages)] == [2, 1, 1]


def test_claude_code_command_is_isolated_and_text_goes_through_stdin(monkeypatch, tmp_path):
    captured = {}
    result = {"is_error": False, "structured_output": FIELDS, "modelUsage": {"claude-haiku-4-5-20251001": {"outputTokens": 40}, "claude-opus-5": {"outputTokens": 90}}}

    def fake_run(command, **kwargs):
        captured.update(command=command, **kwargs)
        return subprocess.CompletedProcess(command, 0, stdout=json.dumps(result), stderr="")

    monkeypatch.setattr(backend_module.subprocess, "run", fake_run)
    secret_text = f"{SYNTHETIC_TEXT} patient {PRIVATE_NAME}"

    classification = ClaudeCodeBackend(model="claude-opus-5").classify(Payload(text=secret_text), tmp_path)

    command = captured["command"]
    assert classification.model == "claude-opus-5"
    assert secret_text in captured["input"]
    assert not any(PRIVATE_NAME in part for part in command)
    for flag in ("--no-session-persistence", "--strict-mcp-config", "--disable-slash-commands", "--json-schema"):
        assert flag in command
    assert command[command.index("--setting-sources") + 1] == ""
    assert command[command.index("--tools") + 1] == ""
    assert captured["env"]["CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC"] == "1"
    vision = ClaudeCodeBackend().command("vision")
    assert vision[vision.index("--tools") + 1] == "Read"


def test_parse_result_errors_never_carry_model_output():
    with pytest.raises(UsageLimitReached):
        parse_result(1, json.dumps({"is_error": True, "result": "Claude usage limit reached"}), "")
    with pytest.raises(BackendError) as failure:
        parse_result(1, json.dumps({"is_error": True, "result": f"cannot read {PRIVATE_NAME}"}), "")
    assert PRIVATE_NAME not in str(failure.value)
    with pytest.raises(BackendError):
        parse_result(0, json.dumps({"is_error": False, "structured_output": {"doc_type": "diagnosis"}}), "")


def test_cli_requires_consent_and_prints_numbers_only(setup, monkeypatch):
    data_dir, _, _, _ = setup
    monkeypatch.setattr(backend_module, "ClaudeCodeBackend", FakeBackend)
    runner = CliRunner()

    refused = runner.invoke(app, ["classify", "--data-dir", str(data_dir)])
    record_consent(data_dir, FakeBackend.name)
    accepted = runner.invoke(app, ["classify", "--data-dir", str(data_dir), "--sample"])

    assert refused.exit_code == 2 and "not confirmed" in refused.output
    assert accepted.exit_code == 0, accepted.output
    assert "This run: classified 5" in accepted.output
    assert "Pages classified: 5 of 7" in accepted.output
    assert PRIVATE_NAME not in accepted.output and "12.03.2019" not in accepted.output


def test_classify_ladder_asks_the_strong_model_only_when_unsure(tmp_path):
    from epicrisis.classify.backend import ModelLadder

    class Unsure(FakeBackend):
        def classify(self, payload, workdir):
            result = super().classify(payload, workdir)
            result.fields["confidence"] = 0.4
            return result

    strong = FakeBackend(model="fake-opus")
    result = ModelLadder(Unsure(model="fake-haiku"), strong).classify(Payload(text="synthetic"), tmp_path)
    assert (len(strong.calls), result.escalation) == (1, ["low_confidence"])

    unused = FakeBackend(model="fake-opus")
    assert ModelLadder(FakeBackend(model="fake-haiku"), unused).classify(Payload(text="synthetic"), tmp_path).escalation is None
    assert unused.calls == []


def test_pages_are_classified_several_at_once_and_all_recorded(setup):
    import threading
    import time

    data_dir, source, _, output = setup

    class Slow(FakeBackend):
        running = peak = 0
        guard = threading.Lock()

        def classify(self, payload, workdir):
            with self.guard:
                Slow.running += 1
                Slow.peak = max(Slow.peak, Slow.running)
            time.sleep(0.05)
            try:
                return super().classify(payload, workdir)
            finally:
                with self.guard:
                    Slow.running -= 1

    stats = classify_source(data_dir, source, Slow(), workers=3)

    assert (stats.classified, Slow.peak) == (7, 3)
    assert len(list(read_records(output / "classify.jsonl"))) == 7
    assert len(list(read_records(output / "ledger.jsonl"))) == 7
    assert classify_source(data_dir, source, FakeBackend(), workers=3).already_done == 7


def test_ladder_keeps_pages_the_small_model_did_alone_unless_they_need_escalation(setup):
    from epicrisis.classify.backend import ModelLadder

    data_dir, source, _, output = setup
    classify_source(data_dir, source, FakeBackend(model="fake-haiku"))
    lines = list(read_records(output / "classify.jsonl"))
    lines[0]["confidence"] = 0.3
    (output / "classify.jsonl").write_text("".join(json.dumps(line) + "\n" for line in lines))

    small, strong = FakeBackend(model="fake-haiku"), FakeBackend(model="fake-opus")
    stats = classify_source(data_dir, source, ModelLadder(small, strong))

    assert (stats.already_done, stats.classified, len(small.calls)) == (6, 1, 1)


def test_pdf_pages_render_safely_from_many_threads(setup):
    from concurrent.futures import ThreadPoolExecutor

    from epicrisis.classify.pages import original_png

    _, _, archive, output = setup
    records = records_by_name(output)
    refs = page_refs(records["scan_a.pdf"]) + page_refs(records["labs_a.pdf"])
    with ThreadPoolExecutor(max_workers=8) as pool:
        images = list(pool.map(lambda ref: original_png(ref, archive), refs * 25))
    assert len(images) == 100 and all(image.startswith(b"\x89PNG") for image in images)
