"""One module per kind of file, and the same four questions asked of each.

Not about any one format — those are tested where they are used, through the walk of a folder and
through the pages a document is made of. This is about the shape: that a reader is a reader.
"""

import inspect

import pytest

from epicrisis.inventory import probes
from epicrisis.readers import READERS, reader_for


def test_every_reader_answers_the_same_questions():
    """probe and pages of every kind; and a page is words or a picture, so at least one of those.

    A format added by copying a branch into two files is how this program came to hold the same
    knowledge twice, in a step that probes and a step that counts pages, with nothing saying what
    the two had to agree about.
    """
    for kind, reader in READERS.items():
        assert callable(getattr(reader, "probe", None)), f"{kind} cannot say what is true of a file"
        assert callable(getattr(reader, "pages", None)), f"{kind} cannot count its pages"
        ways = [name for name in ("text_of", "image_of") if callable(getattr(reader, name, None))]
        assert ways, f"{kind} can produce neither words nor a picture"
        assert list(inspect.signature(reader.probe).parameters) == ["source", "mime"]
        assert list(inspect.signature(reader.pages).parameters) == ["record"]


def test_a_reader_says_the_route_and_the_part_of_every_page_it_counts():
    """A page is "text" or "vision", and belongs to a run of pages inside the file."""
    pages = reader_for("text").pages({"text": {"pages": 3, "of_document": [0, 1, 1]}})

    assert pages == [("text", "text", 0), ("text", "text", 1), ("text", "text", 1)]
    for route, part, _document in pages:
        assert route in ("text", "vision") and isinstance(part, str)


@pytest.mark.parametrize("kind", ["pdf", "image", "word", "excel", "text"])
def test_every_kind_the_walk_can_name_has_a_reader(kind):
    assert reader_for(kind) is not None


def test_a_kind_this_program_does_not_read_has_no_reader_and_is_not_an_error():
    """A file of a kind nobody reads is counted, named and left alone — not a failure."""
    assert reader_for(probes.category("application/x-rar")) is None
    assert reader_for(probes.category("inode/x-empty")) is None
    assert reader_for(None) is None


def test_how_much_of_a_pages_text_a_reader_hands_over_is_one_number_in_one_place(tmp_path):
    """It was three: two MAX_TEXT_CHARS of their own and, in the third, the bare literal.

    Raised for a document and a scan, every sheet of every workbook would have stayed cut where it
    was — and no page of this program says that part of a table never reached the model. Written
    once, in probes.py, where the other measures of a page's text that more than one format asks
    of itself already live, with the reason the ceiling is where it is beside it.
    """
    import io
    import pathlib

    import docx
    import openpyxl

    from epicrisis.readers import excel, pdf, word

    assert (word.MAX_TEXT_CHARS, pdf.MAX_TEXT_CHARS, excel.MAX_TEXT_CHARS) == (probes.MAX_TEXT_CHARS,) * 3
    # And none of the three writes the number out again, under its own name or as a literal.
    spellings = (f"{probes.MAX_TEXT_CHARS}", f"{probes.MAX_TEXT_CHARS:_}", "MAX_TEXT_CHARS =")
    for module in (word, pdf, excel):
        source = pathlib.Path(inspect.getfile(module)).read_text(encoding="utf-8")
        for spelling in spellings:
            assert spelling not in source, f"{module.__name__} says {spelling!r} of its own"

    # The two readers whose page is not a printed page, cutting where the one name says.
    longer_than_a_page = "a" * (probes.MAX_TEXT_CHARS + 500)
    document = docx.Document()
    document.add_paragraph(longer_than_a_page)
    written = io.BytesIO()
    document.save(written)
    assert len(word.text_of(written.getvalue(), None)) == probes.MAX_TEXT_CHARS

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = longer_than_a_page
    book = tmp_path / "one-sheet.xlsx"
    workbook.save(book)

    class Ref:
        index = 0

    assert len(excel.text_of(book.read_bytes(), Ref())) == probes.MAX_TEXT_CHARS
