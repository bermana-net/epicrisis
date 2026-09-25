"""Inventory tests. Every file here is synthetic and generated at test time."""

import json
import shutil
from pathlib import Path

import docx
import openpyxl
import pytest
from PIL import Image
from pypdf import PdfWriter
from typer.testing import CliRunner

from epicrisis.cli import app
from epicrisis.inventory import probes
from epicrisis.inventory.report import Summary, render
from epicrisis.inventory.scan import folder_year_hint, scan, sha256_file

SYNTHETIC_TEXT = (
    "Synthetic hemoglobin 140 g/L reference 120-160 leukocytes 6.1 x10^9/L reference 4.0-9.0 "
    "platelets 250 x10^9/L reference 150-400"
)


def make_text_pdf(path: Path, page_texts: list[str]) -> None:
    """Minimal PDF with a real text layer, one Helvetica line per page."""
    count = len(page_texts)
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids ["
        + b" ".join(b"%d 0 R" % (4 + 2 * i) for i in range(count))
        + b"] /Count %d >>" % count,
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for i, text in enumerate(page_texts):
        content = b"BT /F1 12 Tf 72 720 Td (%s) Tj ET" % text.encode("ascii")
        objects[4 + 2 * i] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i)
        )
        objects[5 + 2 * i] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content)

    data = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(data)
        data += b"%d 0 obj\n%s\nendobj\n" % (number, objects[number])
    xref_offset = len(data)
    size = max(objects) + 1
    data += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for number in range(1, size):
        data += b"%010d 00000 n \n" % offsets[number]
    data += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref_offset)
    path.write_bytes(data)


def make_scan_pdf(path: Path, pages: int = 1) -> None:
    """Image-only PDF, like a scanner produces."""
    images = [Image.new("RGB", (200, 280), "white") for _ in range(pages)]
    images[0].save(path, "PDF", resolution=150, save_all=True, append_images=images[1:])


@pytest.fixture
def archive(tmp_path: Path) -> Path:
    root = tmp_path / "archive"
    (root / "2003").mkdir(parents=True)
    (root / "Lab 2015").mkdir()
    (root / "misc").mkdir()

    make_text_pdf(root / "2003" / "text.pdf", [SYNTHETIC_TEXT, SYNTHETIC_TEXT])
    make_scan_pdf(root / "2003" / "scan.pdf", pages=2)

    scan_part = tmp_path / "scan_part.pdf"
    make_scan_pdf(scan_part)
    writer = PdfWriter()
    writer.append(root / "2003" / "text.pdf")
    writer.append(scan_part)
    writer.write(root / "Lab 2015" / "mixed.pdf")

    Image.new("RGB", (640, 480), "gray").save(root / "Lab 2015" / "photo.jpg", "JPEG", dpi=(300, 300))
    Image.new("RGB", (10, 10), "gray").save(root / "misc" / "looks_like.pdf", "PNG")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in [("analyte", "value"), ("a", 1), ("b", 2)]:
        sheet.append(row)
    workbook.create_sheet("second")["A1"] = "x"
    workbook.save(root / "misc" / "table.xlsx")

    document = docx.Document()
    document.add_paragraph("Synthetic letter")
    document.save(root / "misc" / "letter.docx")

    shutil.copy(root / "2003" / "text.pdf", root / "misc" / "copy.pdf")
    (root / "misc" / "broken.pdf").write_bytes(b"%PDF-1.4\nthis is not a pdf body\n")
    (root / "misc" / ".DS_Store").write_bytes(b"\0" * 16)
    return root


def by_name(records: list[dict]) -> dict[str, dict]:
    return {record["name"]: record for record in records}


def snapshot(root: Path) -> dict:
    return {
        path.relative_to(root).as_posix(): (path.stat().st_size, path.stat().st_mtime_ns, sha256_file(path))
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_records(archive):
    records = by_name(list(scan(archive)))

    text = records["text.pdf"]
    assert text["path"] == "2003/text.pdf"
    assert text["mime"] == "application/pdf"
    assert text["folder_year_hint"] == 2003
    assert text["pdf"]["pages"] == 2
    assert text["pdf"]["text_layer"] == "full"
    assert text["pdf"]["has_images"] is False

    scanned = records["scan.pdf"]["pdf"]
    assert scanned["text_layer"] == "none"
    assert scanned["pages_with_text"] == 0
    assert scanned["has_images"] is True

    mixed = records["mixed.pdf"]
    assert mixed["folder_year_hint"] == 2015
    assert mixed["pdf"]["text_layer"] == "partial"
    assert mixed["pdf"]["pages"] == 3

    photo = records["photo.jpg"]["image"]
    assert (photo["width"], photo["height"], photo["dpi"]) == (640, 480, [300, 300])

    assert records["looks_like.pdf"]["category"] == "image"

    assert records["table.xlsx"]["excel"] == {
        "sheets": 2,
        "sheet_dimensions": [{"rows": 3, "columns": 2}, {"rows": 1, "columns": 1}],
        "embedded_images": 0,
    }
    assert records["letter.docx"]["word"] == {"text_chars": 15, "embedded_images": 0}

    assert "error" in records["broken.pdf"]
    assert records[".DS_Store"]["skipped"] == "os_metadata"
    assert records["copy.pdf"]["sha256"] == text["sha256"]
    assert records["table.xlsx"]["folder_year_hint"] is None


def test_summary(archive):
    summary = Summary()
    for record in scan(archive):
        summary.add(record)

    assert summary.files == 9
    assert summary.skipped == 1
    assert summary.categories["pdf"] == 5
    assert summary.text_layers == {"full": 2, "none": 1, "partial": 1, "damaged": 1}
    assert [sorted(group) for group in summary.duplicate_groups] == [["2003/text.pdf", "misc/copy.pdf"]]
    assert summary.damaged == ["misc/broken.pdf"]
    # scan.pdf 2 + mixed.pdf 1 + photo.jpg 1 + looks_like.pdf 1
    assert summary.vision_pages == 5
    # text.pdf 2 text pages, scan.pdf 2 blank, mixed.pdf 2 text + 1 blank; copy.pdf is a duplicate.
    assert summary.chars_histogram == {"100-499": 4, "0": 3}


def test_duplicates_counted_once_for_vision(archive):
    shutil.copy(archive / "2003" / "scan.pdf", archive / "misc" / "scan_copy.pdf")
    summary = Summary()
    for record in scan(archive):
        summary.add(record)
    assert summary.vision_pages == 5


def test_report_hides_paths_unless_asked(archive):
    summary = Summary()
    for record in scan(archive):
        summary.add(record)

    assert "copy.pdf" not in render(summary)
    assert "misc/copy.pdf" in render(summary, details=True)
    assert "5 x $0.02 = $0.10" in render(summary, usd_per_page=0.02)


def test_cli_writes_jsonl_and_leaves_archive_untouched(archive, tmp_path):
    before = snapshot(archive)
    out = tmp_path / "out" / "inventory.jsonl"

    result = CliRunner().invoke(app, ["inventory", str(archive), "--out", str(out)])

    assert result.exit_code == 0, result.output
    assert "Files: 9" in result.output
    assert "5 x $0.02 = $0.10" in result.output
    assert "PDF pages by visible characters" in result.output
    lines = [json.loads(line) for line in out.read_text(encoding="utf-8").splitlines()]
    assert len(lines) == 10
    assert snapshot(archive) == before


def test_cli_refuses_output_inside_archive(archive):
    before = snapshot(archive)
    result = CliRunner().invoke(app, ["inventory", str(archive), "--out", str(archive / "inv.jsonl")])
    assert result.exit_code == 2
    assert snapshot(archive) == before


def test_stamp_below_threshold_is_not_a_text_layer(tmp_path):
    path = tmp_path / "stamped_scan.pdf"
    make_text_pdf(path, ["SCANNED 1992 PAGE 01", "SCANNED 1992 PAGE 02"])
    facts = probes.probe_pdf(path, "application/pdf")
    assert facts["text_layer"] == "none"
    assert all(0 < chars < probes.MIN_TEXT_CHARS_PER_PAGE for chars in facts["text_chars_per_page"])


def test_garbled_text_layer_counts_as_no_text(tmp_path):
    # A font with a private encoding extracts as control characters, not as the printed text.
    garbled = "".join("\\%03o" % (1 + i % 30) for i in range(150))
    path = tmp_path / "garbled.pdf"
    make_text_pdf(path, [SYNTHETIC_TEXT, garbled])

    facts = probes.probe_pdf(path, "application/pdf")

    assert (facts["garbled_text_pages"], facts["pages_with_text"], facts["text_layer"]) == ([2], 1, "partial")
    assert not probes.text_looks_garbled(SYNTHETIC_TEXT + " Гемоглобін ș ț «»№")


def test_pdf_saved_as_http_response(tmp_path):
    root = tmp_path / "archive"
    root.mkdir()
    pdf = tmp_path / "plain.pdf"
    make_text_pdf(pdf, [SYNTHETIC_TEXT])
    header = b"HTTP/1.0 200 OK\r\nContent-Type: application/pdf\r\nContent-Length: 1\r\n\r\n"
    (root / "download.pdf").write_bytes(header + pdf.read_bytes())
    (root / "chunked.pdf").write_bytes(b"HTTP/1.1 200 OK\r\nTransfer-Encoding: chunked\r\n\r\n5\r\n%PDF-\r\n0\r\n\r\n")

    records = by_name(list(scan(root)))

    wrapped = records["download.pdf"]
    assert wrapped["wrapper"] == {"kind": "http_response", "payload_offset": len(header)}
    assert wrapped["category"] == "pdf"
    assert wrapped["pdf"]["pages"] == 1 and wrapped["pdf"]["text_layer"] == "full"
    assert wrapped["sha256"] == sha256_file(root / "download.pdf")
    assert "unsupported" in records["chunked.pdf"]


@pytest.mark.parametrize(
    "header",
    [
        b"HTTP/1.0 200 OK\n                    Content-Type: application/pdf\n                    Date: Mon\n",
        b"HTTP/1.0 200 OK\r\n  Content-Type: application/pdf\r\n   \r\n",
    ],
    ids=["indented-no-blank-line", "whitespace-blank-line"],
)
def test_http_response_header_variants(tmp_path, header):
    pdf = tmp_path / "plain.pdf"
    make_text_pdf(pdf, [SYNTHETIC_TEXT])
    (tmp_path / "saved.pdf").write_bytes(header + pdf.read_bytes() + b"\n\n" + b"\x00" * 64)

    offset, payload = probes.http_payload(tmp_path / "saved.pdf")

    assert payload.read(5) == b"%PDF-"
    assert offset == len(header)


def test_aes_encrypted_pdf_with_empty_password(tmp_path):
    plain = tmp_path / "plain.pdf"
    make_text_pdf(plain, [SYNTHETIC_TEXT])
    writer = PdfWriter(clone_from=plain)
    writer.encrypt(user_password="", owner_password="synthetic-owner", algorithm="AES-256")
    locked = tmp_path / "locked.pdf"
    writer.write(locked)

    facts = probes.probe_pdf(locked, "application/pdf")

    assert facts["encrypted"] is True
    assert facts["pages"] == 1
    assert facts["text_layer"] == "full"


def test_old_inventory_records_without_image_count(tmp_path):
    summary = Summary()
    summary.add({"path": "a.docx", "size": 1, "folder_year_hint": None, "category": "word", "sha256": "x", "word": {"text_chars": 3}})
    assert summary.office_images == 0


def test_images_inside_office_files_count_for_vision(tmp_path):
    root = tmp_path / "archive"
    root.mkdir()
    picture = tmp_path / "scan.png"
    Image.new("RGB", (40, 40), "white").save(picture)
    document = docx.Document()
    document.add_paragraph("Synthetic cover note")
    document.add_picture(str(picture))
    document.save(root / "pasted.docx")

    record = by_name(list(scan(root)))["pasted.docx"]
    summary = Summary()
    summary.add(record)

    assert record["word"]["embedded_images"] == 1
    assert summary.office_images == 1
    assert summary.vision_pages == 1


def test_damaged_versus_unsupported(tmp_path):
    root = tmp_path / "archive"
    root.mkdir()
    full = tmp_path / "full.jpg"
    Image.new("RGB", (400, 400), "gray").save(full, "JPEG")
    (root / "truncated.jpg").write_bytes(full.read_bytes()[: full.stat().st_size // 2])
    (root / "header_only.jpg").write_bytes(b"\xff\xd8\xff\xe0" + b"\0" * 50)
    (root / "photo.heic").write_bytes(b"\0\0\0\x18ftypheic" + b"\0" * 64)
    (root / "empty.pdf").write_bytes(b"")

    records = by_name(list(scan(root)))
    summary = Summary()
    for record in records.values():
        summary.add(record)

    assert "error" in records["truncated.jpg"]
    assert "error" in records["header_only.jpg"]
    assert "unsupported" in records["photo.heic"]
    assert sorted(summary.damaged) == ["empty.pdf", "header_only.jpg", "truncated.jpg"]
    assert summary.unsupported == ["photo.heic"]
    assert summary.vision_pages == 1


@pytest.mark.parametrize(
    ("relative", "expected"),
    [
        ("2003/file.pdf", 2003),
        ("Lab 2015/sub/file.pdf", 2015),
        ("2010/2012 visit/file.pdf", 2012),
        ("scans/file_2019.pdf", None),
        ("id 123456/file.pdf", None),
        ("2999/file.pdf", None),
    ],
)
def test_folder_year_hint(relative, expected):
    assert folder_year_hint(Path(relative)) == expected


@pytest.mark.parametrize(
    ("mime", "expected"),
    [
        ("application/pdf", "pdf"),
        ("image/tiff", "image"),
        ("application/msword", "legacy_office"),
        ("application/x-ole-storage", "legacy_office"),
        ("inode/x-empty", "empty"),
        ("application/x-rar", "other"),
    ],
)
def test_category(mime, expected):
    assert probes.category(mime) == expected


def test_excel_97_workbooks_are_read_like_new_ones(tmp_path):
    import xlwt

    from epicrisis.classify.pages import materialize, page_refs
    from epicrisis.inventory.scan import inventory_record

    book = xlwt.Workbook()
    sheet = book.add_sheet("Synthetic")
    for row, (name, value) in enumerate([("Synthetic analyte", 71), ("Другий показник", 4.2)]):
        sheet.write(row, 0, name)
        sheet.write(row, 1, value)
    book.save(tmp_path / "old.xls")

    record = inventory_record(tmp_path, tmp_path / "old.xls")

    assert (record["category"], record["excel"]["sheets"], record["excel"]["format"]) == ("excel", 1, "xls")
    [ref] = page_refs(record)
    assert materialize(ref, tmp_path, tmp_path).text == "Synthetic analyte\t71\nДругий показник\t4.2"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_word_97_documents_are_read_through_libreoffice(tmp_path):
    import subprocess

    from epicrisis.classify.pages import materialize, page_refs
    from epicrisis.inventory.scan import inventory_record

    source = tmp_path / "source"
    source.mkdir()
    document = docx.Document()
    document.add_paragraph("Synthetic discharge summary")
    document.save(source / "old.docx")
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{tmp_path / 'profile'}", "--headless", "--convert-to", "doc", "--outdir", str(tmp_path), str(source / "old.docx")],
        check=True, capture_output=True, timeout=120,
    )  # fmt: skip

    record = inventory_record(tmp_path, tmp_path / "old.doc")

    assert (record["category"], record["word"]["format"]) == ("word", "doc")
    [ref] = page_refs(record)
    assert "Synthetic discharge summary" in materialize(ref, tmp_path, tmp_path).text
